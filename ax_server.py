# -*- coding: utf-8 -*-
"""로컬 웹 화면 서버 (표준 라이브러리 http.server — 외부 접속 차단, 내 PC(127.0.0.1)에서만 동작)

  python ax_server.py            → 브라우저에서 http://127.0.0.1:8765 자동 열림
  python ax_server.py --port 9000 --no-browser
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import mimetypes
import sys
import threading
import traceback
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

import ax_db
import ax_match
import ax_rag

WEB_DIR = Path(__file__).resolve().parent / "local_web"
STATE = {"collect": {"running": False, "log": [], "result": None}, "index": None, "index_key": None}
LOCK = threading.Lock()


def _now():
    return dt.datetime.now().isoformat(timespec="seconds")


def get_index(con):
    key = con.execute("SELECT count(*), max(id) FROM chunks").fetchone()
    key = tuple(key)
    if STATE["index_key"] != key:
        STATE["index"] = ax_rag.Index(ax_db.all_chunks(con))
        STATE["index_key"] = key
    return STATE["index"]


def run_collect(opts: dict):
    import ax_collect
    st = STATE["collect"]
    st.update(running=True, log=[], result=None)
    try:
        r = ax_collect.collect(max_pages=int(opts.get("max_pages", 3)), force=bool(opts.get("force")),
                               insecure=bool(opts.get("insecure")), log=lambda m: st["log"].append(m))
        st["result"] = {"new": r["new"], "analyzed": r["analyzed"]}
    except Exception as e:
        st["log"].append(f"! 오류: {e}")
    finally:
        st["running"] = False


def alerts(con, today=None) -> list[dict]:
    """기업별 맞춤 알림: 추천(60점↑) 중 마감 14일 이내 또는 최근 7일 신규 공고"""
    today = today or dt.date.today()
    notices = ax_db.get_notices(con, "접수중")
    recent = (today - dt.timedelta(days=7)).isoformat()
    out = []
    for p in ax_db.list_profiles(con):
        for r in ax_match.match(p, notices, today, include_ineligible=False):
            n = next(x for x in notices if x["project_no"] == r["project_no"])
            is_new = (n.get("first_seen") or "") >= recent
            due = r["d_day"] is not None and r["d_day"] <= 14
            if r["score"] >= 60 and (is_new or due):
                out.append({"profile_id": p["id"], "profile": p.get("name"), "project_no": r["project_no"],
                            "title": r["title"], "score": r["score"], "grade": r["grade"], "d_day": r["d_day"],
                            "kind": ("신규·" if is_new else "") + ("마감임박" if due else "추천"),
                            "detail_url": r["detail_url"]})
    out.sort(key=lambda a: (a["d_day"] if a["d_day"] is not None else 999, -a["score"]))
    return out


def export_matching_xlsx(con) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    today = dt.date.today()
    notices = ax_db.get_notices(con, "접수중")
    wb = Workbook()
    ws = wb.active
    ws.title = "기업별_추천결과"
    head = ["기업명", "순위", "공고번호", "사업명", "적합도", "등급", "마감", "D-day", "지원금액", "지원분야",
            "추천 근거", "주의/부적합 사유", "신청방법", "제출처", "공고 링크"]
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    colors = {"적극 추천": "C6EFCE", "추천": "E2EFDA", "검토": "FFF2CC", "낮음": "F2F2F2", "부적합": "FCE4D6"}
    for p in ax_db.list_profiles(con):
        for i, r in enumerate(ax_match.match(p, notices, today), 1):
            ws.append([p.get("name"), i, r["project_no"], r["title"], r["score"], r["grade"], r["receipt_end"],
                       r["d_day"], r["amount"], ", ".join(r["fields"]),
                       "\n".join(f"(+{x['points']}) {x['text']}" for x in r["reasons"] if x["type"] == "+"),
                       "\n".join(x["text"] for x in r["fails"] + r["warnings"]),
                       ", ".join(r["apply_methods"]), ", ".join(r["emails"]), r["detail_url"]])
            row = ws.max_row
            ws.cell(row, 6).fill = PatternFill("solid", fgColor=colors.get(r["grade"], "FFFFFF"))
            ws.cell(row, 15).hyperlink = r["detail_url"]
    for col, w in zip("ABCDEFGHIJKLMNO", [16, 5, 12, 40, 7, 9, 11, 6, 22, 20, 50, 45, 20, 22, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("공고DB")
    ws2.append(["공고번호", "사업명", "상태", "접수시작", "접수마감", "조기마감", "지원분야", "대상 업종", "대상 기업", "지역",
                "지원금액", "제외조건", "필수조건", "신청방법", "이메일", "요약", "링크"])
    for n in notices:
        t = n.get("target") or {}
        ws2.append([n["project_no"], n["title"], n["status"], n["receipt_start"], n["receipt_end"],
                    "예" if n["early_close"] else "", ", ".join(f["field"] for f in n.get("fields") or []),
                    "·".join(t.get("industries", [])), "·".join(t.get("company_types", [])), "·".join(t.get("regions", [])),
                    ax_match._amount_text(n.get("amount")), "\n".join(x["rule"] for x in t.get("excludes", [])),
                    "\n".join(x["rule"] for x in t.get("requires", [])), ", ".join((n.get("apply") or {}).get("methods", [])),
                    ", ".join((n.get("apply") or {}).get("emails", [])), n.get("summary"), n.get("detail_url")])
    ws3 = wb.create_sheet("알림")
    ws3.append(["기업명", "구분", "사업명", "적합도", "D-day", "링크"])
    for a in alerts(con, today):
        ws3.append([a["profile"], a["kind"], a["title"], a["score"], a["d_day"], a["detail_url"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_profiles_xlsx(con, data: bytes) -> int:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "").strip() for h in rows[0]]
    col = {h: i for i, h in enumerate(head)}

    def g(r, k, default=""):
        i = col.get(k)
        return r[i] if i is not None and i < len(r) and r[i] is not None else default

    def yes(v):
        return str(v).strip() in ("Y", "y", "예", "O", "o", "1", "True", "true", "상장")
    n = 0
    for r in rows[1:]:
        name = str(g(r, "기업명")).strip()
        if not name:
            continue
        p = {"name": name, "biz_no": str(g(r, "사업자번호")).strip(), "region": str(g(r, "소재지(전남/광주/기타)", "전남")).strip(),
             "size": str(g(r, "기업규모", "중소기업")).strip(), "listed": yes(g(r, "상장여부(Y/N)", "N")),
             "industries": [x.strip() for x in str(g(r, "업종(농산,수산,식품,공산품,기타)")).split(",") if x.strip()],
             "items": str(g(r, "주요 수출품목")), "hs_codes": str(g(r, "HS코드")),
             "target_countries": str(g(r, "희망 수출국")), "export_stage": str(g(r, "수출단계", "수출 초보")),
             "export_usd": float(g(r, "전년도 수출실적(USD)", 0) or 0),
             "direct_export": yes(g(r, "직수출(Y/N)", "Y")),
             "is_overseas_distributor": yes(g(r, "해외유통매장운영사(Y/N)", "N")),
             "needs": [x.strip() for x in str(g(r, "필요 지원분야(쉼표구분)")).split(",") if x.strip()]}
        ax_db.save_profile(con, p, _now())
        n += 1
    return n


class Handler(BaseHTTPRequestHandler):
    server_version = "AXMatching/1.0"

    def log_message(self, fmt, *args):
        pass

    def _send(self, code=200, body=b"", ctype="application/json; charset=utf-8", headers=None):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (headers or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _json(self, obj, code=200):
        self._send(code, json.dumps(obj, ensure_ascii=False, default=str).encode("utf-8"))

    def _body(self):
        n = int(self.headers.get("Content-Length") or 0)
        return self.rfile.read(n) if n else b""

    def do_GET(self):
        try:
            self._route("GET")
        except Exception as e:
            self._json({"error": str(e), "trace": traceback.format_exc(limit=3)}, 500)

    def do_POST(self):
        try:
            self._route("POST")
        except Exception as e:
            self._json({"error": str(e), "trace": traceback.format_exc(limit=3)}, 500)

    def do_DELETE(self):
        try:
            self._route("DELETE")
        except Exception as e:
            self._json({"error": str(e)}, 500)

    def _route(self, method):
        u = urlparse(self.path)
        path, qs = u.path, {k: v[0] for k, v in parse_qs(u.query).items()}
        if method == "GET" and path in ("/", "/index.html"):
            return self._send(200, (WEB_DIR / "index.html").read_bytes(), "text/html; charset=utf-8")
        if method == "GET" and path.startswith("/files/"):
            rel = unquote(path[len("/files/"):])
            f = (ax_db.DATA_DIR / rel).resolve()
            if not f.is_relative_to(ax_db.DATA_DIR.resolve()) or not f.is_file():
                return self._json({"error": "not found"}, 404)
            ctype = mimetypes.guess_type(f.name)[0] or "application/octet-stream"
            return self._send(200, f.read_bytes(), ctype,
                              {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(f.name)}"})

        with LOCK:
            con = ax_db.connect()
        try:
            if path == "/api/status" and method == "GET":
                cnt = {r[0]: r[1] for r in con.execute("SELECT status, count(*) FROM notices GROUP BY status")}
                return self._json({"counts": cnt, "last_run": ax_db.last_run(con), "llm": ax_rag.llm_status(),
                                   "collect_running": STATE["collect"]["running"], "fields": ax_match.FIELDS,
                                   "db": str(ax_db.DB_PATH), "today": dt.date.today().isoformat()})
            if path == "/api/notices" and method == "GET":
                return self._json(ax_db.get_notices(con, qs.get("status")))
            if path.startswith("/api/notice/") and method == "GET":
                n = ax_db.get_notice(con, unquote(path.rsplit("/", 1)[1]))
                return self._json(n or {"error": "not found"}, 200 if n else 404)
            if path == "/api/profiles" and method == "GET":
                return self._json(ax_db.list_profiles(con))
            if path == "/api/profiles" and method == "POST":
                pid = ax_db.save_profile(con, json.loads(self._body() or b"{}"), _now())
                con.commit()
                return self._json({"id": pid})
            if path.startswith("/api/profiles/") and method == "DELETE":
                ax_db.delete_profile(con, int(path.rsplit("/", 1)[1]))
                con.commit()
                return self._json({"ok": True})
            if path == "/api/profiles/import" and method == "POST":
                n = import_profiles_xlsx(con, self._body())
                con.commit()
                return self._json({"imported": n})
            if path == "/api/match" and method == "POST":
                p = json.loads(self._body() or b"{}")
                status = p.pop("_status", "접수중")
                return self._json(ax_match.match(p, ax_db.get_notices(con, status)))
            if path == "/api/ask" and method == "POST":
                b = json.loads(self._body() or b"{}")
                notices = {n["project_no"]: n for n in ax_db.get_notices(con)}
                prof = None
                if b.get("profile_id"):
                    prof = next((p for p in ax_db.list_profiles(con) if p["id"] == int(b["profile_id"])), None)
                return self._json(ax_rag.answer(b.get("question", ""), get_index(con), notices,
                                                b.get("project_no") or None, prof, b.get("use_llm", True)))
            if path == "/api/alerts" and method == "GET":
                return self._json(alerts(con))
            if path == "/api/collect" and method == "POST":
                if STATE["collect"]["running"]:
                    return self._json({"error": "이미 수집 중입니다."}, 409)
                opts = json.loads(self._body() or b"{}")
                threading.Thread(target=run_collect, args=(opts,), daemon=True).start()
                return self._json({"started": True})
            if path == "/api/collect/status" and method == "GET":
                st = STATE["collect"]
                return self._json({"running": st["running"], "log": st["log"][-300:], "result": st["result"]})
            if path == "/api/export/matching.xlsx" and method == "GET":
                data = export_matching_xlsx(con)
                name = f"AX_기업별_추천결과_{dt.date.today():%Y%m%d}.xlsx"
                return self._send(200, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                  {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})
            return self._json({"error": f"unknown {method} {path}"}, 404)
        finally:
            con.close()


def main(argv=None):
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--port", type=int, default=8765)
    ap.add_argument("--no-browser", action="store_true")
    a = ap.parse_args(argv)
    ax_db.connect().close()
    srv = ThreadingHTTPServer(("127.0.0.1", a.port), Handler)
    url = f"http://127.0.0.1:{a.port}"
    print(f"AX 수출지원사업 맞춤 매칭 플랫폼 실행 중: {url}  (종료: Ctrl+C)")
    if not a.no_browser:
        threading.Timer(1.0, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("종료합니다.")


if __name__ == "__main__":
    sys.exit(main())
