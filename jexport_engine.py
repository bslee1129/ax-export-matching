# -*- coding: utf-8 -*-
"""
수출e음(jexport.or.kr) '접수중' 지원사업 전체 일괄 다운로드 & 분석기
=====================================================================

동작 순서
  1) 사업공고 목록(/user/reg_biz) 전 페이지를 훑어 상태가 '접수중'인 사업을 모두 수집
  2) 사업별 상세 페이지에서 공고기간·접수기간·담당자·사업개요·지원대상·지원방법·첨부파일 수집
  3) 첨부파일 다운로드 (POST /user/reg_biz/file_download) — ZIP 은 자동 압축해제
  4) 첨부파일 텍스트/표 추출 (HWPX, HWP, PDF, DOCX)
  5) 사업별 분석 리포트 + 전체 종합 리포트(Markdown/Excel/JSON) 생성

설치
  python -m pip install -r requirements.txt
  (requests beautifulsoup4 openpyxl olefile pypdf truststore)

사용 예
  python jexport_open_notices_analyzer.py                    # 접수중 사업 전체
  python jexport_open_notices_analyzer.py --insecure         # SSL 인증서 오류 시
  python jexport_open_notices_analyzer.py --status 전체      # 상태 무관 전체(페이지 수는 --max-pages)
  python jexport_open_notices_analyzer.py --keyword 물류비   # 접수중 중 사업명 필터
  python jexport_open_notices_analyzer.py --skip-download    # 이미 받은 파일로 재분석만
  python jexport_open_notices_analyzer.py --file 공고문.hwp  # 파일 1개만 분석

결과 (기본 폴더: jexport_output\\YYYYMMDD\\)
  00_접수중사업_종합리포트.md / 00_접수중사업_종합.xlsx / 00_접수중사업_종합.json
  <공고번호>_<사업명>\\  첨부파일 원본, *_분석리포트.md, *_분석.xlsx, *_분석결과.json, *_전체텍스트.txt
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import struct
import sys
import time
import zipfile
import zlib
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Optional
from urllib.parse import unquote
import xml.etree.ElementTree as ET

BASE_URL = "https://www.jexport.or.kr"
LIST_URL = BASE_URL + "/user/reg_biz"
LIST_PAGE_URL = BASE_URL + "/user/reg_biz/run"
DETAIL_URL = BASE_URL + "/user/reg_biz/detail"
DOWNLOAD_URL = BASE_URL + "/user/reg_biz/file_download"

HEADERS = {
    "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"),
    "Accept-Language": "ko-KR,ko;q=0.9",
}
REQUEST_DELAY = 0.5   # 서버 부담을 줄이기 위한 요청 간 대기(초)


# ─────────────────────────────────────────────────────────────
# 1. 웹 수집
# ─────────────────────────────────────────────────────────────
@dataclass
class Notice:
    project_no: str
    title: str
    region: str = ""
    period: str = ""
    status: str = ""
    subtitle: str = ""
    detail: dict = field(default_factory=dict)     # 공고기간/접수기간/담당자 등 (th→td)
    sections: dict = field(default_factory=dict)   # 사업개요/지원대상/지원방법 및 제출서류
    attachments: list = field(default_factory=list)
    files: list = field(default_factory=list)      # 저장된(압축해제 포함) 파일 경로
    errors: list = field(default_factory=list)


def _session(insecure: bool = False, ca_bundle: Optional[str] = None):
    """
    SSL 인증서 오류 대응
      - 기본: truststore 로 Windows 인증서 저장소 사용 (브라우저와 같은 방식)
      - --ca-bundle 파일경로 : 직접 지정한 인증서 사용
      - --insecure          : 인증서 검증 끔 (최후 수단)
    """
    if not insecure and not ca_bundle:
        try:
            import truststore
            truststore.inject_into_ssl()
        except ImportError:
            print("  (참고) 'python -m pip install truststore' 를 설치하면 SSL 인증서 오류가 대부분 해결됩니다.")
    import requests
    s = requests.Session()
    s.headers.update(HEADERS)
    if ca_bundle:
        s.verify = ca_bundle
    elif insecure:
        import urllib3
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        s.verify = False
        # 환경변수(REQUESTS_CA_BUNDLE 등)가 session.verify 를 덮어쓰는 requests 동작 방지
        _orig = s.request
        s.request = lambda method, url, **kw: _orig(method, url, **{**kw, "verify": False})
        print("  ! 경고: SSL 인증서 검증을 끄고 접속합니다(--insecure).")
    return s


def _check_ssl_error(e: Exception) -> bool:
    return "CERTIFICATE_VERIFY_FAILED" in str(e) or "SSLError" in type(e).__name__


def _soup(html: str):
    from bs4 import BeautifulSoup
    return BeautifulSoup(html, "html.parser")


def _clean(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def _norm(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


def _safe_name(s: str, maxlen: int = 60) -> str:
    s = re.sub(r'[\\/:*?"<>|★☆\r\n\t]', "", s or "").strip().rstrip(".")
    return s[:maxlen].strip() or "noname"


def _get(sess, url, **kw):
    r = sess.get(url, timeout=30, **kw)
    r.raise_for_status()
    time.sleep(REQUEST_DELAY)
    return r


def parse_list_page(html: str) -> list[Notice]:
    out = []
    for tr in _soup(html).select("tr"):
        a = tr.find("a", href=re.compile(r"f_detail\("))
        if not a:
            continue
        m = re.search(r"f_detail\('([^']+)'\)", a["href"])
        tds = [_clean(td.get_text(" ")) for td in tr.find_all("td")]
        n = Notice(project_no=m.group(1), title=_clean(a.get_text()))
        if len(tds) >= 4:
            n.region, n.period, n.status = tds[0], tds[2], tds[3]
        out.append(n)
    return out


def crawl_notices(sess, max_pages: int = 20, status: str = "접수중") -> list[Notice]:
    """목록을 1페이지부터 순회. status='접수중'이면 접수중이 끝난 뒤 한 페이지 더 확인하고 멈춘다."""
    found: dict[str, Notice] = {}
    pages_without_match = 0
    for page in range(1, max_pages + 1):
        url = LIST_URL if page == 1 else LIST_PAGE_URL
        rows = parse_list_page(_get(sess, url, params={"page": page}).text)
        if not rows:
            break
        matched = 0
        for n in rows:
            if status in ("전체", "all") or status in n.status:
                found.setdefault(n.project_no, n)
                matched += 1
        print(f"  - {page}페이지: {len(rows)}건 중 {matched}건 해당")
        if status not in ("전체", "all"):
            pages_without_match = 0 if matched else pages_without_match + 1
            if pages_without_match >= 2:   # 접수중 공고는 목록 상단에 모여 있음
                break
    return list(found.values())


def fetch_detail(sess, notice: Notice) -> Notice:
    soup = _soup(_get(sess, DETAIL_URL, params={"project_no": notice.project_no, "page": 1}).text)

    h3 = soup.select_one("h3.dcH3Title")
    if h3 and not notice.title:
        notice.title = _clean(h3.get_text())
    st = soup.select_one("p.dcStitle")
    notice.subtitle = _clean(st.get_text()) if st else ""

    info = {}
    for tr in soup.select("tr"):
        cells = tr.find_all(["th", "td"])
        for i, c in enumerate(cells):
            if c.name == "th" and i + 1 < len(cells) and cells[i + 1].name == "td":
                info[_clean(c.get_text())] = _clean(cells[i + 1].get_text(" "))
    notice.detail = info

    for con in soup.select(".dcRegBizCon"):
        h = con.find(["h4", "h5"])
        box = con.select_one(".dcDetailBox") or con
        if h:
            txt = box.get_text("\n")
            txt = "\n".join(l.rstrip() for l in txt.splitlines() if l.strip())
            notice.sections[_clean(h.get_text())] = txt

    for a in soup.find_all("a", onclick=re.compile(r"file_download\(")):
        args = re.findall(r"'([^']*)'", a["onclick"])
        notice.attachments.append({"name": _clean(a.get_text()), "args": args})
    return notice


def _filename_from_cd(cd: str) -> Optional[str]:
    if not cd:
        return None
    m = re.search(r"filename\*\s*=\s*([^']*)''([^;]+)", cd, re.I)
    if m:
        return unquote(m.group(2), encoding=m.group(1) or "utf-8")
    m = re.search(r'filename\s*=\s*"?([^";]+)"?', cd, re.I)
    if not m:
        return None
    raw = m.group(1)
    # 서버가 CP949 바이트를 그대로 보내면 requests 는 latin-1 로 디코딩함 → 복원
    for enc in ("cp949", "utf-8"):
        try:
            return raw.encode("latin-1").decode(enc)
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
    return unquote(raw)


def extract_zip(path: Path) -> list[Path]:
    """ZIP 첨부 압축해제 (한글 파일명 CP949 복원, 경로 조작 방지)"""
    outdir = path.with_suffix("")
    outdir.mkdir(exist_ok=True)
    out = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            name = info.filename
            if not info.flag_bits & 0x800:          # UTF-8 플래그 없으면 CP949 로 재해석
                try:
                    name = name.encode("cp437").decode("cp949")
                except (UnicodeEncodeError, UnicodeDecodeError):
                    pass
            target = outdir / _safe_name(Path(name).name, 120)
            target.write_bytes(z.read(info))
            out.append(target)
            if target.suffix.lower() == ".zip":
                out.extend(extract_zip(target))
    return out


def download_attachments(sess, notice: Notice, outdir: Path) -> list[Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    saved = []
    for att in notice.attachments:
        args = att["args"] or [notice.project_no]
        data = {"project_no": args[0]}
        if len(args) > 1:            # 첨부가 여러 개인 공고 대비
            data["file_no"] = args[1]
        r = sess.post(DOWNLOAD_URL, data=data, timeout=60,
                      headers={"Referer": f"{DETAIL_URL}?project_no={notice.project_no}"})
        r.raise_for_status()
        time.sleep(REQUEST_DELAY)
        if "text/html" in r.headers.get("Content-Type", "") and len(r.content) < 5000:
            notice.errors.append(f"다운로드 실패(HTML 응답): {att['name']}")
            continue
        name = _filename_from_cd(r.headers.get("Content-Disposition", "")) or att["name"] or f"{notice.project_no}.bin"
        path = outdir / _safe_name(name, 120)
        path.write_bytes(r.content)
        print(f"    ✓ 저장: {path.name} ({len(r.content):,} bytes)")
        saved.append(path)
    return saved


def expand_files(paths: list[Path]) -> list[Path]:
    """ZIP 은 풀어서 안의 문서들로 대체"""
    out = []
    for p in paths:
        if p.suffix.lower() == ".zip" or (zipfile.is_zipfile(p) and p.suffix.lower() not in (".hwpx", ".docx", ".xlsx")):
            try:
                inner = extract_zip(p)
                print(f"    ✓ 압축해제: {p.name} → {len(inner)}개 파일")
                out.extend(q for q in inner if q.suffix.lower() != ".zip")
            except zipfile.BadZipFile as e:
                print(f"    ! 압축해제 실패: {p.name} ({e})")
        else:
            out.append(p)
    return out


# ─────────────────────────────────────────────────────────────
# 2. 문서 텍스트 추출
# ─────────────────────────────────────────────────────────────
@dataclass
class Table:
    index: int
    rows: list            # list[list[str]]  (병합셀 위치 반영된 2차원 격자)
    context: str = ""     # 표 바로 앞 문단 (표 제목 추정용)


@dataclass
class Document:
    path: str
    lines: list = field(default_factory=list)    # 본문 문단 (표는 [표 n] 자리표시)
    tables: list = field(default_factory=list)   # list[Table]

    @property
    def text(self) -> str:
        return "\n".join(self.lines)


def _ln(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


class HwpxReader:
    """HWPX(OWPML, zip+xml) 파서: 문단·표(병합셀 포함)·글상자 텍스트를 순서대로 추출"""

    def __init__(self, path: Path):
        self.path = path
        self.doc = Document(str(path))

    def read(self) -> Document:
        with zipfile.ZipFile(self.path) as z:
            sections = sorted(
                (n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n)),
                key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[-1]).group(1)))
            for name in sections:
                root = ET.fromstring(z.read(name))
                self._walk_container(root)
        return self.doc

    # 문단 컨테이너(section, subList 등)
    def _walk_container(self, elem):
        for child in elem:
            tag = _ln(child.tag)
            if tag == "p":
                self._paragraph(child)
            elif tag == "tbl":
                self._table(child)
            else:
                self._walk_container(child)

    def _paragraph(self, p):
        buf = []

        def flush():
            s = "".join(buf).rstrip()
            if s.strip():
                self.doc.lines.append(s)
            buf.clear()

        for run in p:
            if _ln(run.tag) != "run":
                continue
            for item in run:
                tag = _ln(item.tag)
                if tag == "t":
                    buf.append(self._t_text(item))
                elif tag == "tbl":
                    flush()
                    self._table(item)
                elif tag in ("secPr", "ctrl", "linesegarray"):
                    continue
                else:   # 글상자(rect/drawText), 그림 캡션 등 내부 문단
                    flush()
                    self._walk_container(item)
        flush()

    @staticmethod
    def _t_text(t) -> str:
        out = [t.text or ""]
        for c in t:
            tag = _ln(c.tag)
            if tag == "tab":
                out.append("\t")
            elif tag in ("lineBreak",):
                out.append("\n")
            out.append(c.tail or "")
        return "".join(out)

    def _cell_text(self, tc) -> str:
        parts = []
        for el in tc.iter():
            if _ln(el.tag) == "t":
                parts.append(self._t_text(el))
            elif _ln(el.tag) == "p" and parts and not parts[-1].endswith("\n"):
                parts.append("\n")
        txt = "".join(parts)
        return "\n".join(_clean(x) for x in txt.split("\n") if _clean(x))

    def _table(self, tbl):
        cells = []
        max_r = max_c = 0
        for tr in (e for e in tbl if _ln(e.tag) == "tr"):
            for tc in (e for e in tr if _ln(e.tag) == "tc"):
                addr = next((e for e in tc if _ln(e.tag) == "cellAddr"), None)
                span = next((e for e in tc if _ln(e.tag) == "cellSpan"), None)
                r = int(addr.get("rowAddr", 0)) if addr is not None else 0
                c = int(addr.get("colAddr", 0)) if addr is not None else 0
                rs = int(span.get("rowSpan", 1)) if span is not None else 1
                cs = int(span.get("colSpan", 1)) if span is not None else 1
                cells.append((r, c, rs, cs, self._cell_text(tc)))
                max_r, max_c = max(max_r, r + rs), max(max_c, c + cs)
        grid = [["" for _ in range(max_c)] for _ in range(max_r)]
        for r, c, rs, cs, txt in cells:
            for i in range(r, r + rs):          # 병합셀은 같은 값으로 채움(분석 편의)
                for j in range(c, c + cs):
                    grid[i][j] = txt
        ctx = self.doc.lines[-1] if self.doc.lines else ""
        t = Table(index=len(self.doc.tables) + 1, rows=grid, context=ctx)
        self.doc.tables.append(t)
        self.doc.lines.append(f"[표 {t.index}]")


def _hwp_para_text(raw: bytes) -> str:
    """HWPTAG_PARA_TEXT → 문자열. 제어문자: 문자형(1 WCHAR) / 인라인·확장형(8 WCHAR)
    일반 문자는 UTF-16LE 바이트를 모아 한 번에 디코딩(한글 원문자 등 서로게이트 쌍 보존)"""
    buf, k, n = bytearray(), 0, len(raw)
    NL, TAB, SP, HY = "\n".encode("utf-16-le"), "\t".encode("utf-16-le"), " ".encode("utf-16-le"), "-".encode("utf-16-le")
    while k + 1 < n:
        ch = raw[k] | (raw[k + 1] << 8)
        if ch < 32:
            if ch in (10, 13):
                buf += NL; k += 2
            elif ch == 24:
                buf += HY; k += 2
            elif ch in (30, 31):
                buf += SP; k += 2
            elif ch == 0 or 25 <= ch <= 29:
                k += 2
            else:                      # 1~9, 11~12, 14~23 : 8 WCHAR(16바이트) 컨트롤
                if ch == 9:
                    buf += TAB
                k += 16
            continue
        buf += raw[k:k + 2]; k += 2
    return buf.decode("utf-16-le", errors="replace").strip()


def read_hwp(path: Path) -> Document:
    """구형 HWP(HWP 5.x 바이너리) 본문·표 추출 (표는 병합셀 위치까지 복원)"""
    import olefile
    doc = Document(str(path))
    ole = olefile.OleFileIO(str(path))
    try:
        header = ole.openstream("FileHeader").read()
        props = struct.unpack_from("<I", header, 36)[0]
        if props & 0x2:
            raise ValueError("암호가 걸린 HWP 문서입니다.")
        if props & 0x4:
            raise ValueError("배포용 HWP 문서(복사 방지)라 텍스트를 추출할 수 없습니다. 한글에서 PDF로 저장 후 --file 로 분석하세요.")
        compressed = bool(props & 0x1)
        secs = sorted((e for e in ole.listdir() if e[0] == "BodyText" and len(e) == 2),
                      key=lambda e: int(re.sub(r"\D", "", e[1]) or 0))
        TBL_ID = struct.unpack("<I", b" lbt")[0]      # 'tbl ' 컨트롤 ID

        for entry in secs:
            data = ole.openstream(entry).read()
            if compressed:
                data = zlib.decompress(data, -15)
            stack = []   # 열린 표: {level, rows, cols, cells, cur, line_pos}

            def close_table():
                t = stack.pop()
                grid_r = max([c[0] + c[2] for c in t["cells"]] + [t["rows"]])
                grid_c = max([c[1] + c[3] for c in t["cells"]] + [t["cols"]])
                grid = [["" for _ in range(grid_c)] for _ in range(grid_r)]
                for r, c, rs, cs, parts in t["cells"]:
                    txt = "\n".join(x for x in parts if x)
                    for i2 in range(r, min(r + rs, grid_r)):
                        for j2 in range(c, min(c + cs, grid_c)):
                            grid[i2][j2] = txt
                if stack and stack[-1]["cur"] is not None:     # 중첩표 → 바깥 셀 텍스트로 평탄화
                    stack[-1]["cur"][4].append(" / ".join(" | ".join(r) for r in grid))
                    return
                ctx = doc.lines[t["line_pos"] - 1] if t["line_pos"] > 0 else ""
                tb = Table(index=len(doc.tables) + 1, rows=grid, context=ctx)
                doc.tables.append(tb)
                doc.lines.insert(t["line_pos"], f"[표 {tb.index}]")

            i = 0
            while i + 4 <= len(data):
                h = struct.unpack_from("<I", data, i)[0]
                tag, level, size = h & 0x3FF, (h >> 10) & 0x3FF, (h >> 20) & 0xFFF
                i += 4
                if size == 0xFFF:
                    size = struct.unpack_from("<I", data, i)[0]
                    i += 4
                body = data[i:i + size]
                i += size

                while stack and level <= stack[-1]["level"]:
                    close_table()

                if tag == 71 and len(body) >= 4 and struct.unpack_from("<I", body, 0)[0] == TBL_ID:   # CTRL_HEADER 'tbl '
                    stack.append({"level": level, "rows": 0, "cols": 0, "cells": [], "cur": None,
                                  "line_pos": len(doc.lines)})
                elif tag == 77 and stack and level == stack[-1]["level"] + 1 and len(body) >= 8:      # TABLE
                    stack[-1]["rows"], stack[-1]["cols"] = struct.unpack_from("<HH", body, 4)
                elif tag == 72 and stack and stack[-1]["rows"] and level == stack[-1]["level"] + 1 and len(body) >= 16:     # LIST_HEADER(셀)
                    col, row, cs, rs = struct.unpack_from("<HHHH", body, 8)
                    cell = [row, col, max(rs, 1), max(cs, 1), []]
                    stack[-1]["cells"].append(cell)
                    stack[-1]["cur"] = cell
                elif tag == 67:                                                                       # PARA_TEXT
                    s = _hwp_para_text(body)
                    if not s:
                        continue
                    if stack and stack[-1]["cur"] is not None:
                        stack[-1]["cur"][4].append(_clean(s))
                    else:
                        doc.lines.extend(l.rstrip() for l in s.split("\n") if l.strip())
            while stack:
                close_table()
    finally:
        ole.close()
    return doc


def read_pdf(path: Path) -> Document:
    from pypdf import PdfReader
    doc = Document(str(path))
    for page in PdfReader(str(path)).pages:
        doc.lines.extend(l for l in (page.extract_text() or "").splitlines() if l.strip())
    return doc


def read_docx(path: Path) -> Document:
    doc = Document(str(path))
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = next(e for e in root if _ln(e.tag) == "body")
    for el in body:
        if _ln(el.tag) == "p":
            s = "".join(t.text or "" for t in el.iter() if _ln(t.tag) == "t").strip()
            if s:
                doc.lines.append(s)
        elif _ln(el.tag) == "tbl":
            rows = []
            for tr in (e for e in el.iter() if _ln(e.tag) == "tr"):
                rows.append(["".join(t.text or "" for t in tc.iter() if _ln(t.tag) == "t").strip()
                             for tc in tr if _ln(tc.tag) == "tc"])
            t = Table(len(doc.tables) + 1, rows, doc.lines[-1] if doc.lines else "")
            doc.tables.append(t)
            doc.lines.append(f"[표 {t.index}]")
    return doc


def read_xlsx(path: Path) -> Document:
    """엑셀 서식(별첨) → 시트별 표"""
    from openpyxl import load_workbook
    doc = Document(str(path))
    wb = load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        width = max(len(r) for r in rows)
        keep = [j for j in range(width) if any(j < len(r) and r[j] for r in rows)]
        rows = [[r[j] if j < len(r) else "" for j in keep] for r in rows]
        doc.lines.append(f"[시트] {ws.title}")
        t = Table(len(doc.tables) + 1, rows, f"시트: {ws.title}")
        doc.tables.append(t)
        doc.lines.append(f"[표 {t.index}]")
        doc.lines.extend(" ".join(c for c in r if c) for r in rows[:200])
    return doc


def read_document(path: Path) -> Document:
    ext = path.suffix.lower()
    if ext == ".hwpx" or (ext not in (".hwp", ".pdf", ".docx", ".xlsx", ".xlsm") and zipfile.is_zipfile(path)
                          and "Contents/content.hpf" in zipfile.ZipFile(path).namelist()):
        return HwpxReader(path).read()
    if ext == ".hwp":
        return read_hwp(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path)
    if ext == ".zip":
        raise ValueError("ZIP 첨부는 압축을 푼 뒤 --file 로 개별 분석하세요.")
    raise ValueError(f"지원하지 않는 형식: {ext}")


# ─────────────────────────────────────────────────────────────
# 3. 분석
# ─────────────────────────────────────────────────────────────
BULLET = r"[○◦❍∘•●▪■□◆◇※\-‐⁃·]"
KEY_VALUE_RE = re.compile(rf"^\s*{BULLET}?\s*([가-힣A-Za-z()·‧ ]{{2,20}}?)\s*[:：]\s*(.+)$")
HEADING_RES = [
    (1, re.compile(r"^\s*(\d{1,2})\.\s+(\S.*)$")),                 # 1. 수요조사 개요
    (1, re.compile(r"^\s*([󰊱-󰊹])\s*(\S.*)$")),                     # 󰊱 (HWP 원문자)
    (1, re.compile(r"^\s*【\s*(붙임\s*\d+)\s*】\s*(.*)$")),          # 【붙임 1】
    (2, re.compile(r"^\s*([가-하])\.\s+(\S.*)$")),                   # 가. 1단계 ...
]
PHONE_RE = re.compile(r"0\d{1,2}-\d{3,4}-\d{4}")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
MONEY_RE = re.compile(r"(\d[\d,\.]*)\s*(백만\s*원|백만원|천원|만원|억\s*원|억원|원|만\s*달러|달러|불|USD)")
PCT_RE = re.compile(r"(\d{1,3}(?:\.\d+)?)\s*%")
AREA_RE = re.compile(r"(\d+)\s*㎡")
FULL_DATE = re.compile(r"(20\d{2}|'\d{2})\s*[.\-년]\s*(\d{1,2})\s*[.\-월]\s*(\d{1,2})\s*[.일]?")
RANGE_RE = re.compile(
    r"(?P<y1>20\d{2}|'\d{2})\s*[.\-]\s*(?P<m1>\d{1,2})\s*[.\-]\s*(?P<d1>\d{1,2})\.?(?:\s*\([월화수목금토일]\))?"
    r"\s*~\s*"
    r"(?:(?P<y2>20\d{2}|'\d{2})\s*[.\-]\s*)?(?P<m2>\d{1,2})\s*[.\-]\s*(?P<d2>\d{1,2})\.?(?:\s*\([월화수목금토일]\))?"
    r"(?:\s*(?P<hm>\d{1,2}:\d{2}))?")
MONTH_RANGE_RE = re.compile(r"(20\d{2}|'\d{2})\s*\.\s*(\d{1,2})\.?\s*~\s*(?:(20\d{2}|'\d{2})\s*\.\s*)?(\d{1,2})\.?")

TOPIC_KEYS = {
    "사업명": ["사 업 명", "사업명", "조 사 명", "조사명"],
    "사업목적": ["사업목적", "조사목적"],
    "사업기간": ["사업기간", "사업 추진기간"],
    "조사기간": ["조사기간"],
    "지원대상": ["지원대상", "조사대상", "신청대상"],
    "지원내용": ["지원내용"],
    "지원규모": ["지원규모", "지원한도", "지원비율"],
    "운영조건": ["운영조건", "의무 운영기간", "신규개설 기한"],
    "수탁기관": ["수탁기관"],
    "접수기간": ["접수기간", "제출기간"],
    "제출서류": ["제출서류"],
    "신청방법": ["신청방법"],
    "평가방법": ["평가방법", "방법", "선정"],
}


def _year(y: str) -> int:
    return 2000 + int(y[1:]) if y.startswith("'") else int(y)


def normalize_key(k: str) -> str:
    return re.sub(r"\s+", "", k)


def find_headings(lines):
    out = []
    for i, line in enumerate(lines):
        for level, rx in HEADING_RES:
            m = rx.match(line)
            if m and len(line) < 60:
                out.append({"line": i, "level": level, "title": _clean(line)})
                break
    return out


def section_of(headings, idx):
    top = sub = ""
    for h in headings:
        if h["line"] > idx:
            break
        if h["level"] == 1:
            top, sub = h["title"], ""
        else:
            sub = h["title"]
    return top, sub


def extract_key_values(lines, headings):
    kv = []
    for i, line in enumerate(lines):
        m = KEY_VALUE_RE.match(line)
        if not m:
            continue
        key, val = normalize_key(m.group(1)), _clean(m.group(2))
        if not val or len(key) > 12:
            continue
        top, sub = section_of(headings, i)
        # 이어지는 '-' 하위 항목 수집
        subs = []
        for nxt in lines[i + 1:i + 15]:
            s = nxt.strip()
            if re.match(r"^[-‐⁃]\s*", s) and not KEY_VALUE_RE.match(nxt):
                subs.append(re.sub(r"^[-‐⁃]\s*", "", s))
            elif s.startswith("※"):
                subs.append(s)
            else:
                break
        kv.append({"section": top, "subsection": sub, "key": key, "value": val, "items": subs, "line": i})
    return kv


def summarize_topics(kv):
    topics = {}
    for name, keys in TOPIC_KEYS.items():
        wanted = {normalize_key(k) for k in keys}
        hits = [x for x in kv if x["key"] in wanted]
        if hits:
            topics[name] = hits
    return topics


def extract_schedule(lines, headings, today: dt.date):
    """날짜 범위(~) 를 모두 찾아 D-day 계산"""
    out, seen = [], set()
    for i, line in enumerate(lines):
        for m in RANGE_RE.finditer(line):
            y1 = _year(m["y1"])
            y2 = _year(m["y2"]) if m["y2"] else y1
            try:
                start = dt.date(y1, int(m["m1"]), int(m["d1"]))
                end = dt.date(y2, int(m["m2"]), int(m["d2"]))
            except ValueError:
                continue
            if end < start and not m["y2"]:
                end = end.replace(year=end.year + 1)
            label = KEY_VALUE_RE.match(line)
            label = normalize_key(label.group(1)) if label else ""
            top, sub = section_of(headings, i)
            key = (start, end, label, sub)
            if key in seen:
                continue
            seen.add(key)
            d = (end - today).days
            status = "마감" if d < 0 else ("진행중" if start <= today else "예정")
            out.append({"section": top, "subsection": sub, "label": label or sub or top,
                        "start": start.isoformat(), "end": end.isoformat(), "time": m["hm"] or "",
                        "d_day": d, "status": status, "source": _clean(line)})
    # 월 단위 범위 (예: 2026. 12. ~ 2027. 1.)
    for i, line in enumerate(lines):
        if RANGE_RE.search(line):
            continue
        for m in MONTH_RANGE_RE.finditer(line):
            y1 = _year(m.group(1)); y2 = _year(m.group(3)) if m.group(3) else y1
            top, sub = section_of(headings, i)
            label = KEY_VALUE_RE.match(line)
            out.append({"section": top, "subsection": sub,
                        "label": normalize_key(label.group(1)) if label else (sub or top),
                        "start": f"{y1}-{int(m.group(2)):02d}", "end": f"{y2}-{int(m.group(4)):02d}",
                        "time": "", "d_day": None, "status": "월단위", "source": _clean(line)})
    return out


def extract_requirements(lines):
    """'지원조건/자격요건' 뒤에 나열된 조건들을 체크리스트로 추출"""
    reqs, cur = [], None
    for line in lines:
        s = line.strip()
        if re.search(r"(지원조건|자격요건|지원대상\s*:.*다음)", s):
            cur = {"title": _clean(s), "items": []}
            reqs.append(cur)
            continue
        if cur is None:
            continue
        if re.match(r"^([-‐⁃]|\d\))\s*", s) and len(s) < 150:
            cur["items"].append(re.sub(r"^([-‐⁃]|\d\))\s*", "", s))
        elif s.startswith("※") and cur["items"]:
            cur["items"][-1] += f"  ({s})"
        elif cur["items"]:
            cur = None
    return [r for r in reqs if r["items"]]


def extract_numbers(lines):
    money, pct, area = [], [], []
    for line in lines:
        for m in MONEY_RE.finditer(line):
            money.append({"value": m.group(1), "unit": re.sub(r"\s+", "", m.group(2)), "context": _clean(line)[:120]})
        for m in PCT_RE.finditer(line):
            pct.append({"value": float(m.group(1)), "context": _clean(line)[:120]})
        for m in AREA_RE.finditer(line):
            area.append({"value": int(m.group(1)), "context": _clean(line)[:120]})
    return money, pct, area


def extract_attachments_forms(lines):
    forms = []
    for line in lines:
        m = re.search(r"\[?【?\s*붙임\s*(\d+)\s*】?\]?\s*[「]?([^」\n]*?)[」]?\s*(\d+부)?\s*$", line)
        if m and ("붙임" in line) and len(line) < 120 and m.group(2).strip():
            forms.append({"no": int(m.group(1)), "title": _clean(m.group(2).strip(" -「」"))})
    uniq = {}
    for f in forms:
        uniq.setdefault((f["no"], f["title"]), f)
    return sorted(uniq.values(), key=lambda f: f["no"])


def blank_form_fields(doc: Document):
    """서식(신청서/조사서) 표에서 '작성해야 할 항목' 추출: 값이 비어있는 행의 라벨"""
    out = []
    for t in doc.tables:
        if len(t.rows) < 3:
            continue
        fields = []
        for row in t.rows:
            uniq = []
            for c in row:
                if not uniq or uniq[-1] != c:
                    uniq.append(c)
            label = next((c for c in uniq if c), "")
            rest = [c for c in uniq[1:]]
            empty_or_template = any((not c) or re.search(r"(□|:\s*$|:\s*\)|\(\s*\)|월\s*$|㎡|USD|천원)", c) for c in rest) if rest else False
            if label and empty_or_template and len(label) < 40:
                fields.append(" / ".join(x for x in uniq[:2] if x))
        if len(fields) >= 3:
            out.append({"table": t.index, "context": t.context, "fields": fields})
    return out


def analyze(doc: Document, notice: Optional[Notice], today: dt.date) -> dict:
    lines = doc.lines
    headings = find_headings(lines)
    kv = extract_key_values(lines, headings)
    money, pct, area = extract_numbers(lines)
    full = doc.text
    return {
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "today": today.isoformat(),
        "notice": asdict(notice) if notice else None,
        "file": doc.path,
        "stats": {"paragraphs": len(lines), "tables": len(doc.tables), "characters": len(full)},
        "outline": headings,
        "topics": summarize_topics(kv),
        "key_values": kv,
        "schedule": extract_schedule(lines, headings, today),
        "requirements": extract_requirements(lines),
        "forms": extract_attachments_forms(lines),
        "form_fields": blank_form_fields(doc),
        "contacts": {"phones": sorted(set(PHONE_RE.findall(full))),
                     "emails": sorted(set(EMAIL_RE.findall(full)))},
        "money": money, "percent": pct, "area_m2": area,
    }


# ─────────────────────────────────────────────────────────────
# 4. 출력
# ─────────────────────────────────────────────────────────────
def _md_table(rows):
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    esc = lambda s: str(s).replace("|", "\\|").replace("\n", "<br>")
    rows = [list(r) + [""] * (w - len(r)) for r in rows]
    out = ["| " + " | ".join(esc(c) for c in rows[0]) + " |", "|" + "---|" * w]
    out += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows[1:]]
    return "\n".join(out)


def build_report(res: dict, doc: Document) -> str:
    n = res.get("notice") or {}
    L = []
    title = n.get("title") or Path(res["file"]).stem
    L += [f"# 공고 분석 리포트: {title}", "",
          f"- 분석일: {res['today']}",
          f"- 원본 파일: `{Path(res['file']).name}`",
          f"- 문단 {res['stats']['paragraphs']}개 · 표 {res['stats']['tables']}개 · {res['stats']['characters']:,}자"]
    if n:
        L.append(f"- 공고번호: {n.get('project_no')}  / 상태: {n.get('status', '')}  / 접수: {n.get('period', '')}")
        d = n.get("detail") or {}
        for k in ("공고기간", "접수기간", "소속", "이름", "전화번호"):
            if d.get(k):
                L.append(f"- {k}: {d[k]}")
    L.append("")

    # 핵심 요약
    L += ["## 1. 핵심 요약", ""]
    rows = [["항목", "내용", "위치"]]
    for name, hits in res["topics"].items():
        for h in hits[:3]:
            val = h["value"] + ("<br>" + "<br>".join("· " + x for x in h["items"]) if h["items"] else "")
            rows.append([name, val, h["subsection"] or h["section"]])
    L += [_md_table(rows), ""]

    # 일정
    L += ["## 2. 주요 일정 / D-day", ""]
    rows = [["구분", "시작", "종료", "D-day", "상태", "위치"]]
    for s in res["schedule"]:
        dd = "" if s["d_day"] is None else (f"D-{s['d_day']}" if s["d_day"] > 0 else ("D-Day" if s["d_day"] == 0 else f"D+{-s['d_day']}"))
        rows.append([s["label"], s["start"], s["end"] + (f" {s['time']}" if s["time"] else ""), dd, s["status"],
                     s["subsection"] or s["section"]])
    L += [_md_table(rows) if len(rows) > 1 else "_날짜 범위를 찾지 못했습니다._", ""]
    upcoming = [s for s in res["schedule"] if s["d_day"] is not None and s["d_day"] >= 0]
    if upcoming:
        s = min(upcoming, key=lambda x: x["d_day"])
        L += [f"> ⏰ 가장 가까운 마감: **{s['label']}** — {s['end']} {s['time']} (D-{s['d_day']})", ""]

    # 자격요건
    L += ["## 3. 자격요건 체크리스트", ""]
    for r in res["requirements"]:
        L.append(f"**{r['title']}**")
        L += [f"- [ ] {it}" for it in r["items"]]
        L.append("")
    if not res["requirements"]:
        L += ["_자격요건 목록을 찾지 못했습니다._", ""]

    # 제출서류/서식
    L += ["## 4. 제출서류 · 서식", ""]
    L += [f"- 붙임 {f['no']}: {f['title']}" for f in res["forms"]] or ["_없음_"]
    L.append("")
    for ff in res["form_fields"][:5]:
        L.append(f"**작성 항목 (표 {ff['table']}{' · ' + ff['context'][:40] if ff['context'] else ''})**")
        L += [f"- {x}" for x in ff["fields"]]
        L.append("")

    # 금액/비율
    L += ["## 5. 금액 · 비율 · 면적 언급", ""]
    rows = [["값", "단위", "문맥"]] + [[m["value"], m["unit"], m["context"]] for m in res["money"][:40]]
    L += [_md_table(rows), ""]
    rows = [["비율(%)", "문맥"]] + [[p["value"], p["context"]] for p in res["percent"][:30]]
    L += [_md_table(rows), ""]
    if res["area_m2"]:
        L += [_md_table([["면적(㎡)", "문맥"]] + [[a["value"], a["context"]] for a in res["area_m2"]]), ""]

    # 연락처
    L += ["## 6. 문의처 · 제출처", "",
          "- 전화: " + (", ".join(res["contacts"]["phones"]) or "-"),
          "- 이메일: " + (", ".join(res["contacts"]["emails"]) or "-"), ""]

    # 목차
    L += ["## 7. 문서 목차", ""]
    L += [("  " * (h["level"] - 1)) + f"- {h['title']}" for h in res["outline"]]
    L.append("")

    # 표
    L += ["## 8. 문서 내 표", ""]
    for t in doc.tables:
        if not any(any(c for c in r) for r in t.rows):
            continue
        L += [f"### 표 {t.index}" + (f" — {t.context[:60]}" if t.context else ""), "", _md_table(t.rows), ""]
    return "\n".join(L)


def save_excel(res: dict, doc: Document, path: Path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        print("  (openpyxl 미설치 → Excel 생략)")
        return
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="305496")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    ws = wb.active; ws.title = "핵심요약"
    sheet(ws, ["항목", "키", "내용", "세부", "위치"],
          [[n, h["key"], h["value"], "\n".join(h["items"]), h["subsection"] or h["section"]]
           for n, hs in res["topics"].items() for h in hs], [12, 12, 60, 60, 30])
    sheet(wb.create_sheet("일정"), ["구분", "시작", "종료", "시간", "D-day", "상태", "위치", "원문"],
          [[s["label"], s["start"], s["end"], s["time"], s["d_day"], s["status"], s["subsection"] or s["section"], s["source"]]
           for s in res["schedule"]], [18, 12, 12, 8, 8, 8, 30, 70])
    sheet(wb.create_sheet("자격요건"), ["구분", "요건", "충족여부(O/X)", "비고"],
          [[r["title"][:40], it, "", ""] for r in res["requirements"] for it in r["items"]], [30, 80, 14, 30])
    sheet(wb.create_sheet("항목전체"), ["대분류", "중분류", "키", "값", "세부"],
          [[k["section"], k["subsection"], k["key"], k["value"], "\n".join(k["items"])] for k in res["key_values"]],
          [28, 28, 12, 60, 60])
    sheet(wb.create_sheet("금액"), ["값", "단위", "문맥"],
          [[m["value"], m["unit"], m["context"]] for m in res["money"]], [14, 10, 100])
    for t in doc.tables:
        if not any(any(c for c in r) for r in t.rows) or len(t.rows) < 2:
            continue
        ws = wb.create_sheet(f"표{t.index}")
        ws.append([t.context[:200]])
        ws["A1"].font = Font(bold=True, italic=True)
        for r in t.rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_outputs(res: dict, doc: Document, outdir: Path) -> dict:
    stem = re.sub(r'[\\/:*?"<>|★☆]', "", Path(doc.path).stem).strip() or "notice"
    outdir.mkdir(parents=True, exist_ok=True)
    paths = {
        "report": outdir / f"{stem}_분석리포트.md",
        "json": outdir / f"{stem}_분석결과.json",
        "text": outdir / f"{stem}_전체텍스트.txt",
        "excel": outdir / f"{stem}_분석.xlsx",
    }
    paths["report"].write_text(build_report(res, doc), encoding="utf-8-sig", errors="replace")
    paths["json"].write_text(json.dumps(res, ensure_ascii=False, indent=2, default=str), encoding="utf-8", errors="replace")
    txt = doc.text
    for t in doc.tables:
        txt = txt.replace(f"[표 {t.index}]", f"[표 {t.index}]\n" + "\n".join(" | ".join(r) for r in t.rows) + f"\n[/표 {t.index}]", 1)
    paths["text"].write_text(txt, encoding="utf-8-sig", errors="replace")
    save_excel(res, doc, paths["excel"])
    return paths


def print_console_summary(res: dict):
    print("\n" + "=" * 70)
    t = (res.get("notice") or {}).get("title") or Path(res["file"]).name
    print(f" {t}")
    print("=" * 70)
    for name, hits in res["topics"].items():
        print(f" ■ {name:6s}: {hits[0]['value'][:90]}")
    print("-" * 70)
    for s in res["schedule"]:
        if s["d_day"] is not None:
            dd = f"D-{s['d_day']}" if s["d_day"] >= 0 else f"D+{-s['d_day']}"
            print(f" 📅 {s['label'][:18]:18s} {s['start']} ~ {s['end']} {s['time']:5s} [{s['status']} {dd}]")
    print("-" * 70)
    for r in res["requirements"][:1]:
        print(" ✔ " + r["title"][:60])
        for it in r["items"]:
            print("    - " + it[:90])
    print(f" ☎ {', '.join(res['contacts']['phones'])}   ✉ {', '.join(res['contacts']['emails'])}")
    print("=" * 70)



# ─────────────────────────────────────────────────────────────
# 5. 전체 종합 (접수중 사업 목록)
# ─────────────────────────────────────────────────────────────
SUPPORT_RE = re.compile(r"(지원\s*(내용|금액|한도|기준|규모|비율)|사\s*업\s*량|모집\s*규모|지원한도)")
APPLY_RE = re.compile(r"(신청\s*방법|접수|제출|이메일|온라인|문서24)")


def _parse_period(s: str):
    m = re.findall(r"(20\d{2})[-.](\d{1,2})[-.](\d{1,2})", s or "")
    if not m:
        m2 = re.findall(r"(\d{2})\.(\d{2})\.(\d{2})", s or "")
        m = [("20" + y, mo, d) for y, mo, d in m2]
    ds = []
    for y, mo, d in m:
        try:
            ds.append(dt.date(int(y), int(mo), int(d)))
        except ValueError:
            pass
    return (ds[0] if ds else None, ds[-1] if len(ds) > 1 else None)


def _pick_lines(text: str, rx, limit=6, maxlen=160):
    out = []
    lines = [l.strip() for l in (text or "").splitlines() if l.strip()]
    for i, l in enumerate(lines):
        if rx.search(l):
            out.append(_clean(l)[:maxlen])
            # 바로 아래 '-', '※' 세부 항목 1~2줄 포함
            for nxt in lines[i + 1:i + 3]:
                if re.match(r"^[-‐⁃※*]", nxt):
                    out.append("  " + _clean(nxt)[:maxlen])
                else:
                    break
        if len(out) >= limit:
            break
    seen, uniq = set(), []                      # 중복 제거(순서 유지)
    for x in out:
        k = _norm(x)
        if k not in seen:
            seen.add(k); uniq.append(x)
    return uniq


def summarize_notice(n: Notice, doc_results: list, today: dt.date) -> dict:
    recv = n.detail.get("접수기간") or n.period
    start, end = _parse_period(recv)
    d_day = (end - today).days if end else None
    sec = n.sections
    overview = sec.get("사업개요", "")
    target = sec.get("지원대상", "")
    method = next((v for k, v in sec.items() if "지원방법" in k or "제출서류" in k), "")
    page_text = "\n".join(sec.values())

    emails, phones = set(), set()
    for src in [page_text, n.detail.get("이메일", ""), n.detail.get("전화번호", "")]:
        emails |= set(EMAIL_RE.findall(src)); phones |= set(PHONE_RE.findall(src))
    deadlines = []
    for r in doc_results:
        emails |= set(r["contacts"]["emails"]); phones |= set(r["contacts"]["phones"])
        deadlines += [s for s in r["schedule"] if s["d_day"] is not None and s["d_day"] >= 0]
    next_dl = min(deadlines, key=lambda s: s["d_day"]) if deadlines else None

    support = _pick_lines(overview, SUPPORT_RE) or _pick_lines(target + "\n" + overview, SUPPORT_RE)
    if not support:   # 페이지에 없으면 첨부 분석 결과에서
        for r in doc_results:
            for name in ("지원내용", "지원규모"):
                for h in r["topics"].get(name, [])[:2]:
                    support.append(f"{h['key']} : {h['value']}"[:160])
    reqs = []
    for r in doc_results:
        reqs += r["requirements"]
    return {
        "project_no": n.project_no,
        "title": n.title,
        "subtitle": n.subtitle,
        "region": n.region,
        "status": n.status,
        "notice_period": n.detail.get("공고기간", ""),
        "receipt_period": recv,
        "start": start.isoformat() if start else "",
        "end": end.isoformat() if end else "",
        "d_day": d_day,
        "early_close": bool(re.search(r"예산\s*소진", page_text)),
        "target": [_clean(l)[:200] for l in target.splitlines() if l.strip()][:6],
        "support": support[:8],
        "apply": _pick_lines(method, APPLY_RE, limit=8) or [_clean(l)[:160] for l in method.splitlines() if l.strip()][:5],
        "contact": {"dept": n.detail.get("소속", ""), "name": n.detail.get("이름", ""),
                    "phone": n.detail.get("전화번호", ""), "email": n.detail.get("이메일", "")},
        "emails": sorted(emails), "phones": sorted(phones),
        "attachments": [a["name"] for a in n.attachments],
        "analyzed_files": [r["file"] for r in doc_results],
        "reports": [r.get("_report_path", "") for r in doc_results],
        "doc_next_deadline": next_dl,
        "requirements": reqs,
        "sections": sec,
        "errors": n.errors,
    }


def _bullet(x: str) -> str:
    """앞쪽 글머리표(-, ○, o, ◦ 등) 제거"""
    return re.sub(r"^\s*(?:[-‐⁃○◦❍•·▪o]\s+|[-‐⁃○◦❍•·▪](?=\S))+", "", x).strip() or x.strip()


def _dday(d):
    if d is None:
        return "-"
    return "D-Day" if d == 0 else (f"D-{d}" if d > 0 else f"D+{-d}")


def build_summary_report(items: list, today: dt.date, outdir: Path) -> str:
    L = [f"# 수출e음 접수중 지원사업 종합 분석 ({today.isoformat()} 기준)", "",
         f"- 대상: {len(items)}개 사업 (마감 임박 순)",
         f"- 출처: {LIST_URL}", ""]
    L += ["## 1. 한눈에 보기", ""]
    rows = [["#", "사업명", "접수기간", "마감", "지원 내용·한도(요약)", "제출처"]]
    for i, it in enumerate(items, 1):
        sup = " / ".join(_bullet(s) for s in it["support"][:2])[:120]
        rows.append([i, it["title"] + (" ⚠예산소진시 조기마감" if it["early_close"] else ""),
                     it["receipt_period"], _dday(it["d_day"]), sup, ", ".join(it["emails"][:2])])
    L += [_md_table(rows), ""]

    L += ["## 2. 사업별 상세", ""]
    for i, it in enumerate(items, 1):
        L += [f"### {i}. {it['title']}  ({_dday(it['d_day'])})", "",
              f"- 공고번호: {it['project_no']} · 구분: {it['region']} · 접수기간: {it['receipt_period']}"]
        c = it["contact"]
        if any(c.values()):
            L.append(f"- 담당: {c['dept']} {c['name']} ☎ {c['phone']} {c['email']}".rstrip())
        if it["doc_next_deadline"]:
            s = it["doc_next_deadline"]
            L.append(f"- 첨부문서상 가장 가까운 기한: {s['label']} {s['end']} {s['time']} ({_dday(s['d_day'])})")
        for title, key in (("지원대상", "target"), ("지원내용·한도", "support"), ("신청방법·제출", "apply")):
            if it[key]:
                L += ["", f"**{title}**"] + [f"- {_bullet(x)}" for x in it[key]]
        if it["requirements"]:
            L += ["", "**요건 체크리스트 (첨부 공고문에서 추출)**"]
            for r in it["requirements"][:3]:
                L.append(f"- _{_bullet(r['title'])[:70]}_")
                L += [f"  - [ ] {x}" for x in r["items"][:10]]
        L += ["", "**첨부/분석 파일**"]
        L += [f"- 첨부: {a}" for a in it["attachments"]] or ["- 첨부 없음"]
        for rp in it["reports"]:
            if rp:
                rel = Path(rp).relative_to(outdir).as_posix() if Path(rp).is_relative_to(outdir) else rp
                L.append(f"- 분석 리포트: [{Path(rp).name}]({rel.replace(' ', '%20')})")
        for e in it["errors"]:
            L.append(f"- ⚠ {e}")
        L.append("")
    return "\n".join(L)


def save_summary_excel(items: list, path: Path, outdir: Path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        print("  (openpyxl 미설치 → 종합 Excel 생략)")
        return
    wb = Workbook()
    head, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    red = PatternFill("solid", fgColor="FCE4D6")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active; ws.title = "접수중사업"
    sheet(ws, ["공고번호", "사업명", "구분", "접수시작", "접수마감", "D-day", "조기마감", "지원대상", "지원내용·한도",
               "신청방법·제출", "이메일", "전화", "첨부파일", "분석리포트"],
          [[it["project_no"], it["title"], it["region"], it["start"], it["end"], it["d_day"],
            "예산소진시" if it["early_close"] else "", "\n".join(it["target"]), "\n".join(it["support"]),
            "\n".join(it["apply"]), "\n".join(it["emails"]), "\n".join(it["phones"]),
            "\n".join(it["attachments"]), "\n".join(Path(r).name for r in it["reports"] if r)] for it in items],
          [13, 34, 10, 11, 11, 7, 9, 40, 45, 45, 24, 16, 30, 30])
    for row in ws.iter_rows(min_row=2):
        if isinstance(row[5].value, int) and row[5].value <= 14:
            for c in row:
                c.fill = red
        rp = next((r for r in items[row[0].row - 2]["reports"] if r), None)
        if rp:
            try:
                row[13].hyperlink = Path(rp).relative_to(outdir).as_posix()
            except ValueError:
                row[13].hyperlink = rp
            row[13].font = Font(color="0563C1", underline="single")

    sheet(wb.create_sheet("첨부문서_일정"), ["사업명", "구분", "시작", "종료", "시간", "D-day", "상태", "원문"],
          [[it["title"], s["label"], s["start"], s["end"], s["time"], s["d_day"], s["status"], s["source"]]
           for it in items for s in it.get("_schedule", [])], [34, 18, 11, 11, 7, 7, 8, 80])
    sheet(wb.create_sheet("자격요건_체크"), ["사업명", "요건", "충족(O/X)", "비고"],
          [[it["title"], x, "", ""] for it in items for r in it["requirements"] for x in r["items"]],
          [34, 90, 10, 30])
    sheet(wb.create_sheet("페이지원문"), ["사업명", "항목", "내용"],
          [[it["title"], k, v] for it in items for k, v in it["sections"].items()], [34, 18, 120])
    wb.save(path)


def analyze_files(files: list[Path], notice: Optional[Notice], today: dt.date, outdir: Path) -> list:
    results = []
    for f in files:
        try:
            doc = read_document(f)
        except Exception as e:
            msg = f"{f.name}: 분석 불가 ({e})"
            print(f"    ! {msg}")
            if notice:
                notice.errors.append(msg)
            continue
        if not doc.lines:
            msg = f"{f.name}: 텍스트 없음(스캔 이미지 PDF 또는 배포용/암호 문서일 수 있음)"
            print(f"    ! {msg}")
            if notice:
                notice.errors.append(msg)
            continue
        res = analyze(doc, notice, today)
        paths = write_outputs(res, doc, outdir)
        res["_report_path"] = str(paths["report"])
        results.append(res)
        print(f"    ✓ 분석: {f.name} (문단 {res['stats']['paragraphs']} · 표 {res['stats']['tables']} "
              f"· 일정 {len(res['schedule'])} · 요건 {sum(len(r['items']) for r in res['requirements'])})")
    return results


# ─────────────────────────────────────────────────────────────
# 6. main
# ─────────────────────────────────────────────────────────────
DOC_EXT = {".hwp", ".hwpx", ".pdf", ".docx", ".xlsx"}


def main(argv=None):
    for s in (sys.stdout, sys.stderr):     # Windows 콘솔 한글/이모지 깨짐 방지
        try:
            s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    ap = argparse.ArgumentParser(description="수출e음 접수중 지원사업 첨부파일 일괄 다운로드 & 분석")
    ap.add_argument("--status", default="접수중", help="수집할 상태 (접수중 / 접수마감 / 사업종료 / 전체)")
    ap.add_argument("--keyword", help="사업명에 이 단어가 포함된 것만 (공백 무시)")
    ap.add_argument("--project-no", nargs="+", help="공고번호 직접 지정 (여러 개 가능)")
    ap.add_argument("--file", help="첨부파일 1개만 바로 분석 (다운로드 생략)")
    ap.add_argument("--outdir", help="결과 폴더 (기본: jexport_output\\오늘날짜)")
    ap.add_argument("--max-pages", type=int, default=20, help="목록을 훑을 최대 페이지 수")
    ap.add_argument("--skip-download", action="store_true", help="이미 받은 첨부파일이 있으면 다시 받지 않음")
    ap.add_argument("--today", help="D-day 기준일 (YYYY-MM-DD, 기본: 오늘)")
    ap.add_argument("--insecure", action="store_true", help="SSL 인증서 검증 끄기 (인증서 오류 시 최후 수단)")
    ap.add_argument("--ca-bundle", help="사용할 인증서(PEM) 파일 경로")
    args = ap.parse_args(argv)

    today = dt.date.fromisoformat(args.today) if args.today else dt.date.today()
    outdir = Path(args.outdir or Path("jexport_output") / today.strftime("%Y%m%d")).resolve()
    outdir.mkdir(parents=True, exist_ok=True)

    # ── 단일 파일 모드
    if args.file:
        f = Path(args.file)
        files = expand_files([f]) if f.suffix.lower() == ".zip" else [f]
        for r in analyze_files(files, None, today, outdir):
            print_console_summary(r)
        print(f"\n결과 폴더: {outdir}")
        return 0

    sess = _session(args.insecure, args.ca_bundle)
    try:
        if args.project_no:
            notices = [Notice(project_no=p, title="") for p in args.project_no]
        else:
            print(f"[1] 목록 수집 (상태: {args.status})")
            notices = crawl_notices(sess, args.max_pages, args.status)
        if args.keyword:
            notices = [n for n in notices if _norm(args.keyword) in _norm(n.title)]
        if not notices:
            print("  ! 대상 사업이 없습니다.")
            return 1
        print(f"  → 대상 {len(notices)}건")

        items = []
        for i, n in enumerate(notices, 1):
            print(f"\n[{i}/{len(notices)}] {n.title or n.project_no}")
            try:
                fetch_detail(sess, n)
                ndir = outdir / f"{n.project_no}_{_safe_name(n.title, 40)}"
                existing = [p for p in ndir.glob("*") if p.is_file() and p.suffix.lower() in DOC_EXT | {".zip"}] \
                    if args.skip_download and ndir.exists() else []
                if existing:
                    print(f"    · 기존 첨부 사용: {', '.join(p.name for p in existing)}")
                    raw = existing
                elif n.attachments:
                    raw = download_attachments(sess, n, ndir)
                else:
                    raw = []
                    n.errors.append("첨부파일 없음")
                files = [p for p in expand_files(raw) if p.suffix.lower() in DOC_EXT]
                n.files = [str(p) for p in files]
                results = analyze_files(files, n, today, ndir)
            except Exception as e:
                if _check_ssl_error(e):
                    raise
                n.errors.append(f"처리 오류: {type(e).__name__}: {e}")
                print(f"    ! 처리 오류: {e}")
                results = []
            item = summarize_notice(n, results, today)
            item["_schedule"] = [s for r in results for s in r["schedule"]]
            items.append(item)
    except Exception as e:
        if _check_ssl_error(e):
            print("\n  ! SSL 인증서 확인에 실패했습니다. 아래 순서로 시도해 보세요.\n"
                  "    1) python -m pip install truststore   (설치 후 다시 실행)\n"
                  "    2) python jexport_open_notices_analyzer.py --insecure")
            return 2
        raise

    items.sort(key=lambda x: (x["d_day"] is None, x["d_day"] if x["d_day"] is not None else 0))
    tag = "접수중사업" if args.status == "접수중" else f"{_safe_name(args.status)}사업"
    rpt = outdir / f"00_{tag}_종합리포트.md"
    rpt.write_text(build_summary_report(items, today, outdir), encoding="utf-8-sig", errors="replace")
    save_summary_excel(items, outdir / f"00_{tag}_종합.xlsx", outdir)
    (outdir / f"00_{tag}_종합.json").write_text(
        json.dumps(items, ensure_ascii=False, indent=2, default=str), encoding="utf-8", errors="replace")

    print("\n" + "=" * 78)
    print(f" {args.status} 사업 {len(items)}건 ({today.isoformat()} 기준, 마감 임박 순)")
    print("=" * 78)
    for it in items:
        flag = " (예산소진시 조기마감)" if it["early_close"] else ""
        print(f" {_dday(it['d_day']):>6s} | {it['end'] or '-':10s} | {it['title'][:44]}{flag}")
        if it["errors"]:
            print(f"        ⚠ {'; '.join(it['errors'])[:100]}")
    print("=" * 78)
    print(f"결과 폴더: {outdir}")
    print(f"  - 종합 리포트: {rpt.name}")
    print(f"  - 종합 엑셀  : 00_{tag}_종합.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
